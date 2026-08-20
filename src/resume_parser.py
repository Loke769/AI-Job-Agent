from __future__ import annotations
import pathlib
import re
import yaml
from typing import List

from .config import ROOT, get_settings, resolve_path
from .models import ResumeProfile

# Lazy imports for docx / pdf
SKILL_ALIASES = {
    "py": "python",
    "postgres": "sql",
    "postgresql": "sql",
    "mysql": "sql",
    "amazon web services": "aws",
    "amazon s3": "s3",
}

def normalize_text(s: str) -> str:
    return re.sub(r"\s+", " ", s.lower()).strip()

def extract_text_from_docx(path: pathlib.Path) -> str:
    try:
        from docx import Document
        doc = Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs)
        # also tables
        for table in doc.tables:
            for row in table.rows:
                text += "\n" + " | ".join(cell.text for cell in row.cells)
        return text
    except Exception as e:
        print(f"[resume_parser] docx parse failed {path}: {e}")
        return ""

def extract_text_from_pdf(path: pathlib.Path) -> str:
    # try pdfminer / PyPDF2 if available, else empty
    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except ImportError:
        pass
    try:
        from pdfminer.high_level import extract_text as pdf_extract
        return pdf_extract(str(path))
    except ImportError:
        pass
    print(f"[resume_parser] no PDF parser installed for {path}")
    return ""

def extract_text_from_txt(path: pathlib.Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")

def parse_resume_file(path: str | pathlib.Path) -> str:
    p = resolve_path(str(path))
    if not p.exists():
        return ""
    suf = p.suffix.lower()
    if suf == ".docx":
        return extract_text_from_docx(p)
    elif suf == ".pdf":
        return extract_text_from_pdf(p)
    elif suf in (".txt", ".md"):
        return extract_text_from_txt(p)
    else:
        # try docx fallback
        return extract_text_from_docx(p)

def detect_skills(text: str, master_list: List[str]) -> List[str]:
    low = text.lower()
    found = []
    for skill in master_list:
        # word boundary-ish check
        pattern = r"\b" + re.escape(skill.lower()) + r"\b"
        if re.search(pattern, low):
            found.append(skill)
        # alias
        for alias, canonical in SKILL_ALIASES.items():
            if canonical == skill and alias in low and skill not in found:
                found.append(skill)
    return sorted(set(found))

def load_profile() -> ResumeProfile:
    cfg = get_settings()
    profile_cfg = cfg.get("profile", {})
    resume_path = profile_cfg.get("resume_path", "resumes/master_resume.docx")
    parsed_profile = profile_cfg.get("parsed_profile", "resumes/profile.yaml")
    master_list = cfg.get("matching", {}).get("skills_master_list", [])

    # try yaml profile first if exists
    yaml_path = resolve_path(parsed_profile)
    if yaml_path.exists():
        try:
            data = yaml.safe_load(yaml_path.read_text(encoding="utf-8")) or {}
            # ensure raw_text filled from resume if empty
            if not data.get("raw_text"):
                data["raw_text"] = parse_resume_file(resume_path)
            if not data.get("skills") and data.get("raw_text"):
                data["skills"] = detect_skills(data["raw_text"], master_list)
            return ResumeProfile(**data)
        except Exception as e:
            print(f"[resume_parser] yaml profile load failed: {e}")

    # parse docx
    raw_text = parse_resume_file(resume_path)
    if not raw_text:
        # create a demo profile so pipeline still works
        print(f"[resume_parser] no resume found at {resume_path}, using demo profile")
        raw_text = DEMO_RESUME_TEXT
        # also create a demo docx for future runs
        try:
            create_demo_docx(resolve_path(resume_path), raw_text)
        except Exception as e:
            print(f"demo docx creation failed: {e}")

    skills = detect_skills(raw_text, master_list)
    # try to infer name/email from config
    return ResumeProfile(
        name=profile_cfg.get("name", "Alex Morgan"),
        email=profile_cfg.get("email", "alex.morgan@example.com"),
        phone=profile_cfg.get("phone", "+1-555-010-0000"),
        location=profile_cfg.get("location", "Remote, United States"),
        linkedin=profile_cfg.get("linkedin", ""),
        github=profile_cfg.get("github", ""),
        website=profile_cfg.get("website", ""),
        summary=raw_text[:600],
        skills=skills,
        raw_text=raw_text,
        experience=[],
        education=[],
    )

DEMO_RESUME_TEXT = """
Alex Morgan — Data Engineer
Email: alex.morgan@example.com | Phone: +1-555-010-0000 | Remote, United States
LinkedIn: linkedin.com/in/alexmorgan | GitHub: github.com/alexmorgan

SUMMARY
Data Engineer with 4+ years building scalable batch and streaming pipelines on AWS and GCP.
Expert in Python, SQL, Spark/PySpark, Airflow, Databricks, Snowflake, Kafka, Docker & Kubernetes.
Passionate about data quality, cost optimization, and enabling analytics.

SKILLS
Python, SQL, AWS (S3, Glue, Redshift, Lambda), Spark, PySpark, Airflow, Databricks, Snowflake,
Kafka, Docker, Kubernetes, Java, Scala, Linux, Git, Jenkins, dbt, BigQuery, Terraform, Hadoop

EXPERIENCE
Data Engineer — Nova Analytics (2022-Present, Remote)
- Built ELT pipelines processing 2TB/day with Airflow + Spark on AWS, reduced runtime 40%
- Migrated warehouse from on-prem to Snowflake + dbt, improved query performance 3x
- Implemented streaming ingestion with Kafka + S3 + Glue, enabled real-time dashboards
- Containerized jobs with Docker/Kubernetes, CI/CD via Jenkins & Git

Data Engineer — Bright Data Labs (2020-2022)
- Developed PySpark jobs on Databricks for log processing (500M events/day)
- Tuned Redshift and S3 partitioning, cut storage cost 25%
- Collaborated with analytics to define Looker/Tableau models

EDUCATION
B.S. Computer Science — State University (2016-2020)
"""

def create_demo_docx(path: pathlib.Path, text: str):
    from docx import Document
    from docx.shared import Pt
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10)
    for line in text.strip().splitlines():
        p = doc.add_paragraph(line)
        p.style = style
    doc.save(str(path))
    print(f"[resume_parser] demo resume created at {path}")
    # also create yaml profile
    yaml_path = path.parent / "profile.yaml"
    if not yaml_path.exists():
        profile = {
            "name": "Alex Morgan",
            "email": "alex.morgan@example.com",
            "phone": "+1-555-010-0000",
            "location": "Remote, United States",
            "linkedin": "https://linkedin.com/in/alexmorgan",
            "github": "https://github.com/alexmorgan",
            "summary": "Data Engineer with 4+ years building scalable pipelines.",
            "skills": ["python","sql","aws","spark","pyspark","airflow","databricks","snowflake","kafka","docker","kubernetes","java","scala","linux","git","jenkins","redshift","s3","glue","dbt","bigquery","gcp","terraform","hadoop","hive","presto"],
            "raw_text": text,
        }
        yaml_path.write_text(yaml.safe_dump(profile, sort_keys=False), encoding="utf-8")
